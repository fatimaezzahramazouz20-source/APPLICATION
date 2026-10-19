# -*- coding: utf-8 -*-
"""Export Excel des résultats d'analyse."""

from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .anomalies import Anomalie

_TITRE = Font(bold=True, color="FFFFFF")
_FOND = PatternFill("solid", fgColor="1F4E78")


def export_excel(data: dict, computed: dict,
                 anomalies: list[Anomalie]) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        computed["aggregats"].to_excel(xw, sheet_name="Agrégats", index_label="Agrégat")
        computed["ratios"].to_excel(xw, sheet_name="Ratios", index_label="Ratio")
        data["bilan_table"].to_excel(xw, sheet_name="Bilan", index=False)
        data["cpc_table"].to_excel(xw, sheet_name="CPC", index=False)
        pd.DataFrame(
            [{"Gravité": a.gravite, "Exercice": a.exercice, "Anomalie": a.titre,
              "Détail": a.detail, "Valeur": a.valeur, "Seuil": a.seuil}
             for a in anomalies]
        ).to_excel(xw, sheet_name="Anomalies", index=False)

        wb = xw.book
        for name in wb.sheetnames:
            ws = wb[name]
            for cell in ws[1]:
                cell.font = _TITRE
                cell.fill = _FOND
                cell.alignment = Alignment(horizontal="center")
            ws.freeze_panes = "A2"
            # Largeur des colonnes
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col
                             if c.value is not None), default=10)
                letter = get_column_letter(col[0].column)
                ws.column_dimensions[letter].width = min(max(width + 2, 12), 60)
    return buf.getvalue()
